from openpyxl import load_workbook
from openpyxl.chart import LineChart3D,Reference

wb = load_workbook("台積電2019年9月股價.xlsx")
ws = wb.active 
data = Reference(ws, min_col = 2, max_col = 5,
                 min_row = 1, max_row = 19)
labels = Reference(ws, min_col=1, min_row=2, max_row=19)  
chart = LineChart3D()
chart.add_data(data, from_rows=False,
               titles_from_data=True)
chart.set_categories(labels)
chart.title = "台積電2019年9月股價的3D折線圖"
chart.style = 14
chart.x_axis.title = "日期"
chart.y_axis.title = "股價"
ws.add_chart(chart, "I2")
wb.save("台積電2019年9月股價_lineChart3D.xlsx")
wb.close()