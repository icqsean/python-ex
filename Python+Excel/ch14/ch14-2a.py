from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

wb = load_workbook("產品月銷售量.xlsx")
ws = wb.active 
data = Reference(ws,
                 min_col = 2, 
                 min_row = 1,
                 max_row = 7)  
chart = LineChart()
chart.add_data(data,
               titles_from_data=True)
labels = Reference(ws, min_col=1, min_row=2, max_row=7)
chart.set_categories(labels)
chart.title = "網路商店產品銷售量圖表"
chart.x_axis.title = "月份"
chart.y_axis.title = "銷售量"
ws.add_chart(chart, "F2")
wb.save("產品月銷售量圖表2.xlsx")
wb.close()