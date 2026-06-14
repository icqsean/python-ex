from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
 
wb = load_workbook("口味銷售量.xlsx")
ws = wb.active
data = Reference(ws, min_col = 2, min_row = 1, max_row = 5)
labels = Reference(ws, min_col = 1, min_row = 2, max_row = 5)
chart = PieChart()
chart.add_data(data, titles_from_data = True)
chart.set_categories(labels)
chart.title = "口味銷售量的派圖"
ws.add_chart(chart, "D2")
wb.save("口味銷售量_pieChart.xlsx")
wb.close()

