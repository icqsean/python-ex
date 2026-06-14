from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

wb = load_workbook("產品月銷售量.xlsx")
ws = wb.active 
data = Reference(ws,
                 min_col = 2, 
                 max_col = 4,
                 min_row = 1,
                 max_row = 7)
chart = LineChart()
chart.add_data(data,
               titles_from_data=True)
labels = Reference(ws, min_col=1, min_row=2, max_row=7)
chart.set_categories(labels)
chart.title = "各通路產品銷售量圖表"
chart.style = 13
chart.x_axis.title = "月份"
chart.y_axis.title = "銷售量"
ws.add_chart(chart, "F2")
wb.save("產品月銷售量圖表3.xlsx")
wb.close()