from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

wb = load_workbook("成績管理4.xlsx")
ws = wb.active
data = Reference(ws, min_col = 2, max_col = 4,
                 min_row = 1, max_row = 4)
labels = Reference(ws, min_col=1, min_row=2, max_row=4)  
chart = BarChart()
chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)
chart.title = "班上成績的長條圖"
chart.x_axis.title = "學生"
chart.y_axis.title = "成績"
ws.add_chart(chart, "F2")
wb.save("成績管理4_barChart.xlsx")
wb.close()